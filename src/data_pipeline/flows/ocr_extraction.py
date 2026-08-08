import logging
import os
import re
from pathlib import Path
from typing import Any

import fitz
import pytesseract
from docx import Document
from PIL import Image
from prefect import get_run_logger, task
from pypdf import PdfReader

MIN_EXTRACTED_TEXT_CHARS = 50
DEFAULT_PDF_RENDER_SCALE = 2.0


def _get_logger():
    try:
        return get_run_logger()
    except Exception:
        return logging.getLogger(__name__)


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _is_text_sufficient(text: str, min_chars: int) -> bool:
    return len(text.strip()) >= min_chars


def _ocr_pdf_page(doc: fitz.Document, page_index: int, scale: float, lang: str) -> str:
    page = doc.load_page(page_index)
    matrix = fitz.Matrix(scale, scale)
    pixmap = page.get_pixmap(matrix=matrix, alpha=False)
    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
    return pytesseract.image_to_string(image, lang=lang)


def _extract_pdf_pages(source_id: str, file_path: Path) -> list[tuple[str, str]]:
    min_chars = int(os.getenv("OCR_MIN_TEXT_CHARS", str(MIN_EXTRACTED_TEXT_CHARS)).strip() or "50")
    scale = float(
        os.getenv("OCR_PDF_RENDER_SCALE", str(DEFAULT_PDF_RENDER_SCALE)).strip()
        or str(DEFAULT_PDF_RENDER_SCALE)
    )
    lang = os.getenv("OCR_LANG", "pol+eng").strip() or "pol+eng"

    output: list[tuple[str, str]] = []
    reader = PdfReader(str(file_path))

    with fitz.open(str(file_path)) as doc:
        for page_no, page in enumerate(reader.pages, start=1):
            native_text = _normalize_text(page.extract_text() or "")
            if _is_text_sufficient(native_text, min_chars):
                final_text = native_text
            else:
                ocr_text = _normalize_text(_ocr_pdf_page(doc, page_no - 1, scale, lang))
                final_text = ocr_text or native_text

            if not final_text:
                continue

            output.append((f"{source_id}#page={page_no}", final_text))

    return output


def _extract_docx_text(file_path: Path) -> str:
    document = Document(str(file_path))
    parts: list[str] = []

    parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_values = [
                cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
            ]
            if row_values:
                parts.append(" | ".join(row_values))

    return _normalize_text("\n".join(parts))


def _legacy_passthrough(acquired: Any) -> list[tuple[str, str]]:
    if isinstance(acquired, str):
        return [("synthetic://default", acquired)] if acquired.strip() else []

    if isinstance(acquired, list) and acquired:
        if isinstance(acquired[0], str):
            return [(f"legacy://{i}", p) for i, p in enumerate(acquired) if str(p).strip()]
        if isinstance(acquired[0], tuple) and len(acquired[0]) == 2:
            return [
                (str(source_id), str(content))
                for source_id, content in acquired
                if str(content).strip()
            ]

    return []


@task
def ocr_extraction(acquired_documents: Any) -> list[tuple[str, str]]:
    logger = _get_logger()

    if not acquired_documents:
        return []

    if not isinstance(acquired_documents, list) or (
        acquired_documents and not isinstance(acquired_documents[0], dict)
    ):
        return _legacy_passthrough(acquired_documents)

    tesseract_cmd = os.getenv("TESSERACT_CMD", "").strip()
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    extracted: list[tuple[str, str]] = []

    for item in acquired_documents:
        source_id = str(item.get("source_id", "")).strip()
        path_raw = str(item.get("path", "")).strip()

        if not source_id or not path_raw:
            logger.warning("Skipping invalid acquisition item: %s", item)
            continue

        file_path = Path(path_raw)
        if not file_path.exists() or not file_path.is_file():
            logger.warning("Skipping missing file for source_id=%s path=%s", source_id, file_path)
            continue

        suffix = file_path.suffix.lower()

        try:
            if suffix == ".pdf":
                extracted.extend(_extract_pdf_pages(source_id, file_path))
            elif suffix in {".txt", ".md"}:
                text = _normalize_text(file_path.read_text(encoding="utf-8", errors="ignore"))
                if text:
                    extracted.append((source_id, text))
            elif suffix == ".docx":
                text = _extract_docx_text(file_path)
                if text:
                    extracted.append((source_id, text))
            else:
                logger.debug("Skipping unsupported extension during extraction: %s", file_path)
        except Exception as exc:
            logger.warning("Extraction failed for source_id=%s (%s): %s", source_id, file_path, exc)

    return extracted
